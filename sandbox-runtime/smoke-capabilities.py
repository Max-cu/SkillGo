import json, os, subprocess, tempfile
from pathlib import Path
import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from PIL import Image
import pandas as pd
results={}
with tempfile.TemporaryDirectory() as temp:
 root=Path(temp)
 document=Document(); document.add_paragraph('Capability smoke 中文'); document.save(root/'sample.docx')
 workbook=Workbook(); workbook.active['A1']='中文'; workbook.save(root/'sample.xlsx')
 presentation=Presentation(); presentation.slides.add_slide(presentation.slide_layouts[6]); presentation.save(root/'sample.pptx')
 results['office_create']=all((root/name).stat().st_size>0 for name in ['sample.docx','sample.xlsx','sample.pptx'])
 converted=subprocess.run(['libreoffice', '-env:UserInstallation=file://'+temp+'/lo-profile', '--headless','--convert-to','pdf','--outdir',temp,str(root/'sample.docx')],capture_output=True,timeout=45)
 assert converted.returncode==0 and (root/'sample.pdf').exists(), converted.stderr.decode(errors='replace')
 pdf=pymupdf.open(root/'sample.pdf'); assert len(pdf)>0; assert pdf[0].get_pixmap().width>0; results['office_convert_pdf_render']=True; pdf.close()
 checked=subprocess.run(['qpdf','--check',str(root/'sample.pdf')],capture_output=True,timeout=10); assert checked.returncode==0; results['qpdf']=True
 doc=pymupdf.open(); page=doc.new_page()
 font=subprocess.check_output(['fc-match','-f','%{file}','Noto Sans CJK SC'],text=True).strip()
 page.insert_text((40,80),'中文字体测试',fontname='cjk',fontfile=font,color=(0,0,1))
 assert '中文' in page.get_text(); page.get_pixmap().save(root/'cjk.png'); results['cjk_text_and_render']=True
 assert Image.open(root/'cjk.png').width>0
 assert pd.DataFrame({'x':[1,2]}).x.sum()==3; results['data_tabular']=True
 media=subprocess.run(['ffmpeg','-v','error','-f','lavfi','-i','color=c=white:s=16x16','-frames:v','1','-threads','1',str(root/'media.png')],capture_output=True,timeout=10)
 assert media.returncode==0 and Image.open(root/'media.png').size==(16,16); results['media_convert']=True
print(json.dumps(results))
